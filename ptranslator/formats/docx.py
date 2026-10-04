"""DOCX: translate paragraphs (body, tables, headers, footers), keep the
first run's formatting, and mark paragraphs as right-to-left."""
from __future__ import annotations

import io

from .base import Document, is_translatable


def _iter_paragraphs(container):
    for p in container.paragraphs:
        yield p
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def load_docx(data: bytes) -> Document:
    try:
        import docx  # python-docx
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError as e:
        raise RuntimeError("DOCX support needs: pip install python-docx") from e

    d = docx.Document(io.BytesIO(data))
    paras, seen = [], set()

    def add(container):
        for p in _iter_paragraphs(container):
            if id(p._p) in seen:
                continue          # merged table cells repeat the same paragraph
            seen.add(id(p._p))
            paras.append(p)

    add(d)
    for s in d.sections:
        for part in (s.header, s.footer):
            try:
                if not part.is_linked_to_previous:
                    add(part)
            except Exception:
                pass

    segments, mapping = [], []
    for p in paras:
        texts = p._p.xpath(".//w:t")
        full = "".join(t.text or "" for t in texts)
        if is_translatable(full):
            mapping.append((p, len(segments)))
            segments.append(full.strip())

    def _set_rtl(p):
        pPr = p._p.get_or_add_pPr()
        if pPr.find(qn("w:bidi")) is None:
            pPr.append(OxmlElement("w:bidi"))
        for r in p._p.xpath(".//w:r"):
            rPr = r.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                r.insert(0, rPr)
            if rPr.find(qn("w:rtl")) is None:
                rPr.append(OxmlElement("w:rtl"))

    def build(tr: dict) -> bytes:
        for p, sid in mapping:
            if sid not in tr:
                continue
            texts = p._p.xpath(".//w:t")
            if not texts:
                continue
            texts[0].text = tr[sid]
            texts[0].set(qn("xml:space"), "preserve")
            for t in texts[1:]:
                t.text = ""
            _set_rtl(p)
        out = io.BytesIO()
        d.save(out)
        return out.getvalue()

    return Document(segments, build, ".docx")
