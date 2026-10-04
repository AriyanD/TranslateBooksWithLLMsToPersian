"""PDF: pull the text out page by page, glue broken lines back into
paragraphs, and write the Persian result as a right-to-left DOCX
(one page of the PDF = one page in the DOCX).

Writing Persian back *into* a PDF would break the letters (Persian needs
shaping + RTL), so the output is a clean .docx you can open in Word /
LibreOffice / Google Docs and "Save as PDF" in one click if you want.
Scanned PDFs (pictures of pages, no real text) need OCR first."""
from __future__ import annotations

import io
import re
from collections import Counter

from .base import Document, is_translatable, split_long

MAX_PARAGRAPH = 4000
_END = re.compile(r"[.!?…:;\"'»”)\]]\s*$")


def _clean(line: str) -> str:
    return re.sub(r"[ \t\u00a0]+", " ", line).strip()


def _paragraphs(page_text: str) -> list[str]:
    lines = [_clean(l) for l in page_text.replace("\r", "\n").split("\n")]
    full = [len(l) for l in lines if l]
    typical = sorted(full)[int(len(full) * 0.8)] if full else 0
    paras, cur = [], ""
    for i, line in enumerate(lines):
        if not line:                      # blank line = paragraph break
            if cur:
                paras.append(cur)
            cur = ""
            continue
        if cur:
            if cur.endswith("-") and not cur.endswith(" -") and line[:1].islower():
                cur = cur[:-1] + line     # de-hyphenate "transla-\ntion"
            else:
                cur = cur + " " + line
        else:
            cur = line
        # a short line that ends a sentence is the end of a paragraph
        short = typical and len(line) < typical * 0.75
        nxt = next((l for l in lines[i + 1:] if l), "")
        if short and (_END.search(line) or (nxt[:1].isupper() and len(line) < typical * 0.5)):
            paras.append(cur)
            cur = ""
    if cur:
        paras.append(cur)
    return paras


def _extract_pages(data: bytes) -> list[str]:
    try:
        from pypdf import PdfReader
    except ImportError:
        PdfReader = None
    if PdfReader is not None:
        r = PdfReader(io.BytesIO(data))
        if r.is_encrypted:
            try:
                r.decrypt("")
            except Exception as e:
                raise RuntimeError("This PDF is password protected.") from e
        return [(p.extract_text() or "") for p in r.pages]
    try:
        import pdfplumber
    except ImportError as e:
        raise RuntimeError("PDF support needs: pip install pypdf") from e
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return [(p.extract_text() or "") for p in pdf.pages]


def _drop_running_headers(pages: list[list[str]]) -> list[list[str]]:
    """Remove a header/footer line repeated on most pages (book title etc.)."""
    if len(pages) < 4:
        return pages
    edges = Counter()
    for ps in pages:
        for p in {*(ps[:1]), *(ps[-1:])}:
            if len(p) < 80:
                edges[re.sub(r"\d+", "#", p)] += 1
    common = {k for k, n in edges.items() if n >= len(pages) * 0.5}
    if not common:
        return pages
    out = []
    for ps in pages:
        ps = list(ps)
        if ps and re.sub(r"\d+", "#", ps[0]) in common:
            ps.pop(0)
        if ps and re.sub(r"\d+", "#", ps[-1]) in common:
            ps.pop()
        out.append(ps)
    return out


def load_pdf(data: bytes) -> Document:
    try:
        import docx
        from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as e:
        raise RuntimeError("PDF support needs: pip install python-docx pypdf") from e

    pages = [_paragraphs(t) for t in _extract_pages(data)]
    pages = _drop_running_headers(pages)
    if not any(is_translatable(p) for ps in pages for p in ps):
        raise ValueError("No text found in this PDF. It is probably a scanned book (images of pages). "
                         "Run OCR on it first (e.g. ocrmypdf, or open it in Google Docs) and try again.")

    segments: list[str] = []
    layout: list[list[tuple[str, list[int]]]] = []   # per page: (kind, ids-or-text)
    for ps in pages:
        items = []
        for p in ps:
            if not is_translatable(p):
                items.append(("raw", p))
                continue
            ids = []
            for piece in split_long(p, MAX_PARAGRAPH):
                ids.append(len(segments))
                segments.append(piece)
            items.append(("seg", ids))
        layout.append(items)

    def _rtl(par):
        pPr = par._p.get_or_add_pPr()
        if pPr.find(qn("w:bidi")) is None:
            pPr.append(OxmlElement("w:bidi"))
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in par.runs:
            rPr = r._r.get_or_add_rPr()
            if rPr.find(qn("w:rtl")) is None:
                rPr.append(OxmlElement("w:rtl"))
            fonts = rPr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rPr.insert(0, fonts)
            fonts.set(qn("w:cs"), "Vazirmatn")

    def build(tr: dict) -> bytes:
        d = docx.Document()
        st = d.styles["Normal"]
        st.font.size = docx.shared.Pt(12)
        first = True
        for items in layout:
            if not first:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first = False
            for kind, val in items:
                text = val if kind == "raw" else " ".join(tr.get(i, segments[i]) for i in val)
                par = d.add_paragraph(text)
                _rtl(par)
        out = io.BytesIO()
        d.save(out)
        return out.getvalue()

    return Document(segments, build, ".docx")
